"""Phase 13: safe CSV/Excel import, academic calendar, and report branding."""
from __future__ import annotations

import csv
import io
import sqlite3
from datetime import date
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook
from PIL import Image as PILImage, UnidentifiedImageError
from pydantic import BaseModel, Field, field_validator

MAX_FILE_BYTES = 3 * 1024 * 1024
MAX_ROWS = 1500
EXPECTED = ("id", "name", "age", "email", "course")
EVENT_TYPES = ("exam", "assignment", "holiday", "meeting", "event")


def setup_phase13_tables(conn):
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS calendar_events (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        title TEXT NOT NULL, event_date TEXT NOT NULL,
        event_type TEXT NOT NULL, description TEXT NOT NULL DEFAULT '',
        course TEXT NOT NULL DEFAULT '', semester TEXT NOT NULL DEFAULT '',
        created_by INTEGER REFERENCES users(id) ON DELETE SET NULL
    );
    CREATE INDEX IF NOT EXISTS idx_calendar_events_date ON calendar_events(event_date);
    CREATE TABLE IF NOT EXISTS report_branding (
        id INTEGER PRIMARY KEY CHECK (id=1),
        institution_name TEXT NOT NULL,
        department TEXT NOT NULL,
        contact TEXT NOT NULL,
        footer TEXT NOT NULL,
        logo_png BLOB
    );
    INSERT OR IGNORE INTO report_branding (id,institution_name,department,contact,footer)
    VALUES (1,'CampusFlow Academic Workspace','Office of Academic Records','','Confidential academic report');
    """)


class CalendarEvent(BaseModel):
    title: str = Field(min_length=2, max_length=120)
    event_date: date
    event_type: str = Field(default="event")
    description: str = Field(default="", max_length=1200)
    course: str = Field(default="", max_length=120)
    semester: str = Field(default="", max_length=60)

    @field_validator("title", "description", "course", "semester")
    @classmethod
    def clean(cls, v):
        return v.strip()

    @field_validator("event_type")
    @classmethod
    def valid_type(cls, v):
        if v not in EVENT_TYPES:
            raise ValueError("Invalid academic event type")
        return v


class BrandingUpdate(BaseModel):
    institution_name: str = Field(min_length=2, max_length=110)
    department: str = Field(default="", max_length=110)
    contact: str = Field(default="", max_length=140)
    footer: str = Field(default="", max_length=140)

    @field_validator("institution_name", "department", "contact", "footer")
    @classmethod
    def strip(cls, text):
        return text.strip()


def public_branding(conn):
    row = conn.execute("SELECT institution_name,department,contact,footer, logo_png IS NOT NULL AS has_logo FROM report_branding WHERE id=1").fetchone()
    return dict(row)


def branding_for_pdf(conn):
    row = conn.execute("SELECT institution_name,department,contact,footer,logo_png FROM report_branding WHERE id=1").fetchone()
    return dict(row)


def sanitized_logo(raw: bytes):
    if not raw or len(raw) > 1_500_000:
        raise HTTPException(422, "Upload a PNG or JPEG logo no larger than 1.5 MB")
    try:
        with PILImage.open(io.BytesIO(raw)) as im:
            if im.format not in {"PNG", "JPEG"} or im.width > 2048 or im.height > 2048:
                raise HTTPException(422, "Use a PNG/JPEG image of at most 2048 × 2048 pixels")
            im.load()
            out = io.BytesIO()
            im.convert("RGBA").save(out, "PNG", optimize=True)
            result = out.getvalue()
            if len(result) > 1_500_000:
                raise HTTPException(422, "The processed image is too large")
            return result
    except (OSError, ValueError, UnidentifiedImageError) as exc:
        raise HTTPException(422, "Invalid image file") from exc


def parse_import_file(name: str, raw: bytes, student_model) -> dict:
    """Pure validation/normalization. Caller handles authorization and atomic commit."""
    if not raw or len(raw) > MAX_FILE_BYTES:
        raise HTTPException(422, "Choose a nonempty CSV/XLSX file smaller than 3 MB")
    filename = name.lower()
    try:
        if filename.endswith('.csv'):
            text = raw.decode('utf-8-sig')
            reader = csv.DictReader(io.StringIO(text, newline=''))
            if not reader.fieldnames:
                raise HTTPException(422, 'The CSV has no header row')
            columns = [str(h).strip().lower() for h in reader.fieldnames]
            if len(set(columns)) != len(columns) or set(columns) != set(EXPECTED):
                raise HTTPException(422, 'Required columns: id, name, age, email, course')
            records = [dict(zip(columns, [item.get(header) for header in reader.fieldnames])) for item in reader]
        elif filename.endswith('.xlsx'):
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            sheet = book.active
            items = sheet.iter_rows(values_only=True)
            head = next(items, None)
            columns = [str(h).strip().lower() if h is not None else '' for h in (head or [])]
            if len(set(columns)) != len(columns) or set(columns) != set(EXPECTED):
                raise HTTPException(422, 'Required columns: id, name, age, email, course')
            records = [dict(zip(columns, values)) for values in items if any(v is not None and str(v).strip() for v in values)]
            book.close()
        else:
            raise HTTPException(422, 'Only .csv and .xlsx files are supported')
    except (UnicodeError, csv.Error, ValueError, OSError, EOFError) as exc:
        raise HTTPException(422, 'Could not read the file; use UTF-8 CSV or valid XLSX') from exc
    if len(records) > MAX_ROWS:
        raise HTTPException(422, f'Files can contain at most {MAX_ROWS} students')
    if not records:
        raise HTTPException(422, 'No student rows found')

    errors = []
    valid = []
    seen = set()
    for index, raw_row in enumerate(records, start=2):
        try:
            from pydantic import ValidationError
            for key, value in raw_row.items():
                if isinstance(value, str) and value.strip().startswith('='):
                    raise ValueError('Spreadsheet formulas are not accepted')
            row = student_model.model_validate(raw_row)
            if row.id in seen:
                raise ValueError('Duplicate student ID inside the uploaded file')
            seen.add(row.id)
            valid.append(row.model_dump())
        except (ValidationError, ValueError, TypeError) as exc:
            errors.append({'row': index, 'message': str(exc).split('\n')[0][:160]})
    return {'valid': valid, 'errors': errors, 'total': len(records)}


def preview_import(conn, parsed):
    ids = [r['id'] for r in parsed['valid']]
    existing = set()
    for start in range(0, len(ids), 400):
        part = ids[start:start + 400]
        if part:
            existing.update(r[0] for r in conn.execute('SELECT id FROM students WHERE id IN (' + ','.join('?'*len(part)) + ')',part))
    ready = [r for r in parsed['valid'] if r['id'] not in existing]
    return {
        'total_rows': parsed['total'], 'ready_count': len(ready),
        'existing_count': len(parsed['valid']) - len(ready),
        'invalid_count': len(parsed['errors']), 'errors': parsed['errors'][:12],
        'preview': [{**row, 'status': 'existing' if row['id'] in existing else 'ready'} for row in parsed['valid'][:10]],
        'ready_rows': ready,
    }
